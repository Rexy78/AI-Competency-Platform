import re
import fitz  # PyMuPDF
import docx
from typing import Dict, Any, List
from app.ai.taxonomy import SKILL_TAXONOMY, categorize_skills

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            text += page.get_text("text") + "\n"
    except Exception as e:
        text = file_bytes.decode("utf-8", errors="ignore")
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    import io
    text = ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    except Exception as e:
        text = file_bytes.decode("utf-8", errors="ignore")
    return text

def extract_personal_info(text: str) -> Dict[str, Any]:
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    phone_pattern = r'\(?\+?[0-9]{1,4}\)?[-. ]?\(?[0-9]{2,4}\)?[-. ]?[0-9]{3,4}[-. ]?[0-9]{3,4}'
    linkedin_pattern = r'linkedin\.com/in/[a-zA-Z0-9_-]+'
    github_pattern = r'github\.com/[a-zA-Z0-9_-]+'

    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)
    linkedins = re.findall(linkedin_pattern, text, re.IGNORECASE)
    githubs = re.findall(github_pattern, text, re.IGNORECASE)

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    name = ""
    title = ""

    if lines:
        for line in lines[:5]:
            if not re.search(email_pattern, line) and not re.search(phone_pattern, line) and len(line.split()) <= 5:
                if not any(header in line.lower() for header in ["resume", "curriculum", "page", "profile", "summary"]):
                    name = line
                    break
        for line in lines[1:8]:
            if any(t in line.lower() for t in ["engineer", "developer", "architect", "analyst", "manager", "specialist", "intern", "student"]):
                title = line
                break

    return {
        "name": name if name else (lines[0] if lines else ""),
        "title": title,
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
        "location": "",
        "linkedin": f"linkedin.com/in/{linkedins[0]}" if linkedins else "",
        "github": f"github.com/{githubs[0]}" if githubs else ""
    }

def extract_skills_from_text(text: str) -> List[str]:
    found_skills = set()
    lowered_text = text.lower()
    
    for category, skill_list in SKILL_TAXONOMY.items():
        for skill in skill_list:
            escaped = re.escape(skill.lower())
            if skill.lower() in ["c", "r", "go"]:
                pattern = r'\b' + escaped + r'\b'
            else:
                pattern = r'(?:\b|_)' + escaped + r'(?:\b|_)'
                
            if re.search(pattern, lowered_text):
                found_skills.add(skill)

    return list(found_skills)

def extract_sections(text: str) -> Dict[str, Any]:
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    
    sections = {
        "summary": "",
        "education": [],
        "experience": [],
        "projects": [],
        "certifications": [],
        "achievements": []
    }
    
    current_section = "summary"
    summary_lines = []
    exp_blocks = []
    edu_blocks = []
    proj_blocks = []
    
    for line in lines:
        lower_line = line.lower()
        if any(h in lower_line for h in ["experience", "work history", "employment history", "professional experience", "work experience"]):
            current_section = "experience"
            continue
        elif any(h in lower_line for h in ["education", "academic qualification", "academic background", "educational background"]):
            current_section = "education"
            continue
        elif any(h in lower_line for h in ["projects", "key projects", "academic projects", "personal projects"]):
            current_section = "projects"
            continue
        elif any(h in lower_line for h in ["certifications", "certificates"]):
            current_section = "certifications"
            continue
        elif any(h in lower_line for h in ["summary", "profile", "about me", "professional summary", "career objective"]):
            current_section = "summary"
            continue
            
        if current_section == "summary" and len(summary_lines) < 6:
            summary_lines.append(line)
        elif current_section == "experience":
            exp_blocks.append(line)
        elif current_section == "education":
            edu_blocks.append(line)
        elif current_section == "projects":
            proj_blocks.append(line)

    sections["summary"] = " ".join(summary_lines) if summary_lines else ""
    
    # Parse experience lines cleanly without placeholder text
    formatted_exp = []
    if exp_blocks:
        i = 0
        while i < len(exp_blocks):
            title_line = exp_blocks[i]
            desc_lines = []
            i += 1
            while i < len(exp_blocks) and not any(k in exp_blocks[i].lower() for k in ["engineer", "developer", "architect", "lead", "intern", "analyst", "manager", "201", "202"]):
                desc_lines.append(exp_blocks[i])
                i += 1
                if len(desc_lines) >= 4:
                    break
            formatted_exp.append({
                "title": title_line,
                "company": "",
                "duration": "",
                "description": " ".join(desc_lines) if desc_lines else title_line
            })
    sections["experience"] = formatted_exp

    # Parse education lines cleanly without placeholder text
    formatted_edu = []
    if edu_blocks:
        for line in edu_blocks:
            formatted_edu.append({
                "degree": line,
                "institution": "",
                "year": ""
            })
    sections["education"] = formatted_edu

    # Parse project lines cleanly without placeholder text
    formatted_proj = []
    if proj_blocks:
        for line in proj_blocks:
            formatted_proj.append({
                "title": line,
                "tech": "",
                "description": line
            })
    sections["projects"] = formatted_proj

    return sections

def parse_resume_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    if filename.endswith(".pdf"):
        raw_text = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx") or filename.endswith(".doc"):
        raw_text = extract_text_from_docx(file_bytes)
    else:
        raw_text = file_bytes.decode("utf-8", errors="ignore")

    personal_info = extract_personal_info(raw_text)
    all_skills = extract_skills_from_text(raw_text)
    categorized_skills = categorize_skills(all_skills)
    sections = extract_sections(raw_text)

    return {
        "personal_info": personal_info,
        "sections": sections,
        "skills_categorized": categorized_skills,
        "all_skills": all_skills,
        "raw_text": raw_text
    }
